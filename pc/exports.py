"""The record sheet (shared/PROTOCOL.md, «Record sheet»): one worksheet per class of the quiz
(right-to-left), row 1 = ID, اسم الطالب, #, 1..N, the sections in their column order, المجموع; one row
per active student in class-list order with the points earned per question, the sections as =SUM()
formulas and المجموع as =C<r>.
"""

from __future__ import annotations

import io
import re
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter

import grading

FONT = "Arial"
WIDTHS = {"A": 10.63, "B": 25.63, "C": 5.63}
OTHER_WIDTH = 13
BAD_TITLE = re.compile(r"[\[\]:*?/\\]")


def class_students(store, class_id: str) -> list[dict]:
    """Active students of a class in class-list order. An archived class keeps the students who were active in it
    when it was archived (archivedActive), so its old record sheets keep their rows."""
    cls = store.get("classes", class_id)
    if cls and cls.get("archived"):
        rows = [s for s in store.all("students") if class_id in (s.get("classIds") or []) and s.get("archivedActive")]
    else:
        rows = [s for s in store.all("students") if class_id in (s.get("classIds") or []) and s.get("active", True)
                and not s.get("archived")]
    return sorted(rows, key=lambda s: (s.get("order") or 9999, s.get("name", "")))


def quiz_classes(store, quiz: dict) -> list[dict]:
    classes = {c["id"]: c for c in store.all("classes")}
    ids = [c for c in quiz.get("classIds") or [] if c in classes]
    if not ids:   # no class chosen: every class that has a scanned student
        sids = {s.get("studentId") for s in store.all("scans") if s.get("quizId") == quiz["id"]}
        ids = sorted({c for st in store.all("students") if st["id"] in sids for c in st.get("classIds") or [] if c in classes},
                     key=lambda c: classes[c].get("prefix") or 0)
    return [classes[c] for c in ids]


def _sheet_title(name: str, used: set) -> str:
    base = BAD_TITLE.sub("-", name).strip()[:31] or "Sheet"
    title, k = base, 2
    while title in used:
        title = f"{base[:28]} {k}"
        k += 1
    used.add(title)
    return title


def record_workbook(store, quiz: dict) -> Workbook:
    n = int(quiz.get("questionCount") or 0)
    sections = [s for s in quiz.get("sections") or [] if 1 <= s.get("from", 0) <= s.get("to", 0) <= n]
    scans = grading.counted_scans([s for s in store.all("scans") if s.get("quizId") == quiz["id"]])
    by_student = {s["studentId"]: s for s in scans if s.get("studentId")}
    q_col = lambda i: get_column_letter(4 + i)   # noqa: E731  question i (0-based) -> column
    head = ["ID", "اسم الطالب", "#"] + list(range(1, n + 1)) + [s["name"] for s in sections] + ["المجموع"]
    bold_cols = {1, 2} | set(range(4 + n, 4 + n + len(sections) + 1))
    wb = Workbook()
    wb.remove(wb.active)
    used: set = set()
    classes = quiz_classes(store, quiz) or [{"id": "", "name": "Sheet"}]
    for cls in classes:
        ws = wb.create_sheet(_sheet_title(cls["name"], used))
        ws.sheet_view.rightToLeft = True
        ws.append(head)
        for st in class_students(store, cls["id"]) if cls["id"] else []:
            r = ws.max_row + 1
            scan = by_student.get(st["id"])
            row = [int(st["gradalId"]) if str(st.get("gradalId", "")).isdigit() else st.get("gradalId"), st["name"]]
            if scan:
                key = grading.key_for(quiz, scan.get("version"))
                answers = scan.get("answers") or []
                per_q = []
                for i in range(n):
                    k = key[i] if i < len(key) else ""
                    a = answers[i] if i < len(answers) else ""
                    per_q.append(grading.points_for(quiz, i) if grading.is_correct(k, a) else 0)
                row += [scan.get("earned", 0)] + per_q
            else:
                row += [None] * (1 + n)
            row += [f"=SUM({q_col(s['from'] - 1)}{r}:{q_col(s['to'] - 1)}{r})" for s in sections]
            row.append(f"=C{r}")
            ws.append(row)
        for row in ws.iter_rows():
            for cell in row:
                cell.font = Font(name=FONT, size=10, bold=cell.row == 1 and cell.column in bold_cols)
                cell.alignment = Alignment(horizontal="center", vertical="center")
        for col in range(1, len(head) + 1):
            letter = get_column_letter(col)
            ws.column_dimensions[letter].width = WIDTHS.get(letter, OTHER_WIDTH)
        ws.row_dimensions[1].height = 22.5
        for r in range(2, ws.max_row + 1):
            ws.row_dimensions[r].height = 17.5
    return wb


def record_bytes(store, quiz: dict) -> bytes:
    buf = io.BytesIO()
    record_workbook(store, quiz).save(buf)
    return buf.getvalue()


def record_path(folder: Path, quiz: dict) -> Path:
    name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "-", quiz.get("name") or "اختبار").strip(" .") or "اختبار"
    return Path(folder) / f"{name} ({quiz.get('code')}).xlsx"


def save_record(store, quiz: dict, folder: Path) -> Path:
    """Writes the record sheet next to the others. Raises PermissionError when Excel holds it open."""
    path = record_path(folder, quiz)
    path.parent.mkdir(parents=True, exist_ok=True)
    data = record_bytes(store, quiz)
    tmp = path.with_suffix(".tmp")
    tmp.write_bytes(data)
    try:
        tmp.replace(path)
    except PermissionError:
        tmp.unlink(missing_ok=True)
        raise
    return path
