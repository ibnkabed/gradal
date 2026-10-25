"""Gradal PC tests: the sheet (both directions) read back by the reference reader, grading, class lists,
the question maker's checks, and the HTTP API the window and the phone use.

    python -m unittest discover -s pc/tests -v
"""

from __future__ import annotations

import io
import json
import random
import sys
import tempfile
import threading
import unittest
import urllib.error
import urllib.request
from pathlib import Path

PC = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(PC))
sys.path.insert(0, str(PC.parent / "shared"))

from openpyxl import Workbook, load_workbook  # noqa: E402

import build_layouts  # noqa: E402
import exam_maker  # noqa: E402
import grading  # noqa: E402
import omr  # noqa: E402
import roster  # noqa: E402
import server  # noqa: E402
import sheet_layout  # noqa: E402
import sheets  # noqa: E402
import students_import  # noqa: E402

SETTINGS = {"schoolLines": ["المملكة العربية السعودية", "وزارة التعليم"]}


def quiz(n, language, key=None, sections=None, code=4407):
    q = {"code": code, "name": "Test", "language": language, "questionCount": n, "sections": sections or [],
         "header": {"subject": "Science", "title": "Quiz"}}
    q["sheet"] = sheet_layout.compute(q, sheets.load_layout())
    q["keys"] = [key or [""] * n]
    return q


class SheetTests(unittest.TestCase):
    def test_encode_matches_layout_builder(self):
        L = sheets.load_layout()
        for code in (1000, 4407, 9999):
            self.assertEqual(sheets.encode(code, 0, L), build_layouts.encode(code, 0))

    def test_key_sheet_reads_back_in_both_directions(self):
        L = sheets.load_layout()
        rng = random.Random(7)
        for lang, n in (("ar", 20), ("en", 33), ("ar", 60), ("en", 5)):
            key = [rng.choice("ABCD") for _ in range(n)]
            q = quiz(n, lang, key, [{"name": "S1", "from": 1, "to": n // 2}, {"name": "S2", "from": n // 2 + 1, "to": n}])
            shapes = sheets.page_shapes(L, quiz=q, settings=SETTINGS, key=key, gradal_id="")
            read = omr.read_sheet(sheets.page_raster(shapes, 6.0), L, n, question_bubbles=q["sheet"]["questionBubbles"])
            self.assertIsNotNone(read, (lang, n))
            self.assertEqual((read.quiz_code, read.answers), (4407, key), (lang, n))

    def test_personalised_id(self):
        L = sheets.load_layout()
        q = quiz(20, "ar")
        shapes = sheets.page_shapes(L, quiz=q, settings=SETTINGS, student_name="أحمد", gradal_id="312", class_name="أول 1")
        read = omr.read_sheet(sheets.page_raster(shapes, 6.0), L, 20, question_bubbles=q["sheet"]["questionBubbles"])
        self.assertEqual((read.student_id, read.answers), ("312", [""] * 20))

    def test_arabic_mirrors_choices(self):
        L = sheets.load_layout()
        ar, en = quiz(10, "ar")["sheet"], quiz(10, "en")["sheet"]
        self.assertGreater(ar["questionBubbles"][0][0][0], ar["questionBubbles"][0][3][0])   # أ is rightmost
        self.assertLess(en["questionBubbles"][0][0][0], en["questionBubbles"][0][3][0])
        svg = sheets.page_svg(sheets.page_shapes(L, quiz=quiz(10, "ar"), settings=SETTINGS))
        for text in ("رقم الطالب", "صحيح", "‏الاسم‏"):
            self.assertIn(text, svg)
        self.assertNotIn("Student ID", svg)

    def test_limits(self):
        L = sheets.load_layout()
        for n in (4, 61):
            with self.assertRaises(sheet_layout.LayoutError):
                sheet_layout.compute({"questionCount": n}, L)


class GradingTests(unittest.TestCase):
    def test_grade(self):
        q = {"questionCount": 4, "keys": [["A", "BC", "", "D"]], "points": [1, 2, 1, 0.5]}
        self.assertEqual(grading.grade(q, ["A", "C", "A", "AD"], 0), (3, 3.5))


class ImportTests(unittest.TestCase):
    def test_xlsx_with_header_and_class(self):
        wb = Workbook()
        ws = wb.active
        ws.append(["تقرير الطلاب"])
        ws.append(["م", "رقم الهوية", "اسم الطالب", "الفصل"])
        ws.append([1, "1133582856", "أحمد محمد العتيبي", "أول متوسط 1"])
        ws.append([2, "1133582857", "خالد سعد القحطاني", "أول متوسط 2"])
        buf = io.BytesIO()
        wb.save(buf)
        out = students_import.read_file(buf.getvalue(), "noor.xlsx")
        self.assertEqual([c["name"] for c in out["classes"]], ["أول متوسط 1", "أول متوسط 2"])
        self.assertEqual(out["classes"][0]["students"][0], {"name": "أحمد محمد العتيبي", "nationalId": "1133582856"})

    def test_html_table_saved_as_xls(self):
        page = ("<html><body><table><tr><th>الاسم</th></tr><tr><td>فهد علي الشهري</td></tr>"
                "<tr><td>سعد ناصر الدوسري</td></tr></table></body></html>").encode("utf-8")
        out = students_import.read_file(page, "list.xls", "ثاني متوسط 3")
        self.assertEqual([s["name"] for s in out["classes"][0]["students"]], ["فهد علي الشهري", "سعد ناصر الدوسري"])

    def test_paste_and_merge(self):
        from store import Store
        st = Store(Path(tempfile.mkdtemp()))
        roster.merge(st, students_import.read_pasted("1. أحمد\n2. خالد\n\n", "أول 1"))
        roster.merge(st, students_import.read_pasted("خالد\nسعد", "أول 1"))
        ids = {s["name"]: (s["gradalId"], s["active"]) for s in st.all("students")}
        self.assertEqual(ids, {"أحمد": ("101", False), "خالد": ("102", True), "سعد": ("103", True)})


class MakerTests(unittest.TestCase):
    def test_clean_item(self):
        ok = {"type": "mcq", "text": "ما عاصمة المملكة؟", "options": ["الرياض", "جدة", "مكة", "الدمام"], "answer": "0"}
        self.assertEqual(exam_maker.clean_item(ok)["answer"], 0)
        for bad in ({**ok, "options": ["أ", "أ", "ب", "ج"]}, {**ok, "options": ["الرياض", "جدة", "مكة", "جميع ما سبق"]},
                    {**ok, "answer": 4}, {"type": "tf", "text": "عبارة صحيحة هنا", "answer": "ربما"}):
            self.assertIsNone(exam_maker.clean_item(bad))
        self.assertEqual(exam_maker.clean_item({"type": "tf", "text": "الشمس نجم.", "answer": "صح"})["answer"], True)

    def test_balance(self):
        items = [{"type": "mcq", "text": str(i), "options": ["r", "a", "b", "c"], "answer": 0} for i in range(20)]
        exam_maker.balance(items, random.Random(3))
        letters = [q["answer"] for q in items]
        self.assertEqual(sorted(letters), sorted([i % 4 for i in range(20)]))
        self.assertTrue(all(not (letters[i] == letters[i - 1] == letters[i - 2]) for i in range(2, 20)))
        self.assertTrue(all(q["options"][q["answer"]] == "r" for q in items))

    def test_parse_json_with_fences(self):
        self.assertEqual(exam_maker.parse_json('```json\n{"questions": []}\n```'), {"questions": []})


class ApiTests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.dir = Path(tempfile.mkdtemp())
        cls.srv = server.make_server(cls.dir, "127.0.0.1", 0)
        cls.port = cls.srv.server_address[1]
        threading.Thread(target=cls.srv.serve_forever, daemon=True).start()

    @classmethod
    def tearDownClass(cls):
        cls.srv.shutdown()

    def call(self, method, path, body=None, token=None, raw=False, headers=None):
        h = {"X-Gradal": "1", **(headers or {})}
        data = None
        if body is not None:
            data = body if isinstance(body, bytes) else json.dumps(body).encode("utf-8")
            h.setdefault("Content-Type", "application/json")
        if token:
            h["Authorization"] = f"Bearer {token}"
        req = urllib.request.Request(f"http://127.0.0.1:{self.port}{path}", data=data, method=method, headers=h)
        try:
            with urllib.request.urlopen(req, timeout=120) as r:
                out = r.read()
                return r.status, (out if raw else json.loads(out or b"{}"))
        except urllib.error.HTTPError as e:
            return e.code, json.loads(e.read() or b"{}")

    def pair(self):
        self.call("POST", "/m/pair/start", {"deviceName": "iPhone"})
        pin = self.call("GET", "/api/state")[1]["pin"]["code"]
        return self.call("POST", "/m/pair", {"pin": pin, "deviceId": "dev-1", "deviceName": "iPhone"})[1]["token"]

    def test_round_trip(self):
        self.assertEqual(self.call("GET", "/hello")[1]["app"], "gradal-public")
        self.call("POST", "/api/students/paste", {"text": "أحمد محمد\nخالد سعد", "className": "أول 1"})
        st = self.call("GET", "/api/state")[1]
        cid = st["classes"][0]["id"]
        code, r = self.call("POST", "/api/quizzes", {"name": "قصير", "language": "ar", "classIds": [cid],
                                                       "questionCount": 10, "keys": [list("ABCDABCDAB")]})
        self.assertEqual(code, 200)
        q = r["quiz"]
        token = self.pair()
        snap = self.call("GET", "/m/snapshot", token=token)[1]
        self.assertTrue(any(x["code"] == q["code"] and x["language"] == "ar" for x in snap["quizzes"]))
        scans = [{"id": "bbbbbbbb-0000-4000-8000-000000000001", "quizId": q["id"], "gradalId": "101",
                  "answers": list("ABCDABCDAA"), "scannedAt": "2026-10-21T08:30:00Z", "updatedAt": "2026-10-21T08:30:00Z"}]
        self.assertEqual(self.call("POST", "/m/scans", {"scans": scans}, token=token)[1]["accepted"], [scans[0]["id"]])
        res = self.call("GET", f"/api/quizzes/{q['id']}/results")[1]
        self.assertEqual((res["rows"][0]["studentName"], res["rows"][0]["earned"]), ("أحمد محمد", 9))
        files = self.call("GET", f"/m/prints?quizId={q['id']}", token=token)[1]["files"]
        self.assertEqual([f["id"] for f in files], ["key", f"class-{cid}", "blank"])
        code, pdf = self.call("GET", f"/m/print?quizId={q['id']}&id=blank", token=token, raw=True)
        self.assertTrue(pdf.startswith(b"%PDF"))
        code, xlsx = self.call("GET", f"/api/quizzes/{q['id']}/record", raw=True)
        ws = load_workbook(io.BytesIO(xlsx)).worksheets[0]
        self.assertEqual((ws["A2"].value, ws["C2"].value), (101, 9))

    def test_paper_approval(self):
        st = self.call("GET", "/api/state")[1]
        app = self.srv.app
        qs = [{"type": "mcq", "text": f"سؤال رقم {i}؟", "options": ["أ1", "ب1", "ج1", "د1"], "answer": i % 4} for i in range(6)]
        qs += [{"type": "tf", "text": "عبارة صحيحة.", "answer": True}, {"type": "tf", "text": "عبارة خاطئة.", "answer": False}]
        pid = "11111111-2222-4333-8444-555555555555"
        app.store.data["papers"][pid] = {"id": pid, "status": "draft", "createdAt": "2026-10-25T00:00:00Z", "questions": qs,
                                         "request": {"name": "مسودة", "subject": "العلوم", "grade": "", "language": "ar",
                                                     "classIds": [], "notes": ""}}
        quiz = self.call("POST", f"/api/papers/{pid}/approve", {})[1]["quiz"]
        self.assertEqual(quiz["keys"][0], ["A", "B", "C", "D", "A", "B", "A", "B"])
        self.assertEqual([s["name"] for s in quiz["sections"]], ["اختيار من متعدد", "صح أو خطأ"])
        bad = self.call("POST", f"/api/papers/{pid}", {"questions": [{"type": "mcq", "text": "ناقص؟", "options": ["1"], "answer": 0}]})
        self.assertEqual(bad[0], 400)

    def test_guards(self):
        self.assertEqual(self.call("GET", "/m/snapshot")[0], 401)
        req = urllib.request.Request(f"http://127.0.0.1:{self.port}/api/settings", method="POST", data=b"{}",
                                     headers={"Content-Type": "application/json"})
        with self.assertRaises(urllib.error.HTTPError) as cm:
            urllib.request.urlopen(req, timeout=10)
        self.assertEqual(cm.exception.code, 403)
        self.assertFalse(any("key" in p for p in self.call("GET", "/api/state")[1]["ai"]))   # never the key itself


if __name__ == "__main__":
    unittest.main()
