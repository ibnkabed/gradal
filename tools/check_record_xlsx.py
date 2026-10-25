"""Checks the sample record sheet written by the iPhone app's own .xlsx writer.

    python tools/check_record_xlsx.py <record.xlsx>

The file comes from `OMRCheck --xlsx <out.xlsx>` (ios/Gradal/Record/RecordSheet.swift, sample()):
two classes, 17 questions, the Midterm sections. Checks it against shared/PROTOCOL.md
«Record sheet». Exit code 0 only when everything matches.
"""

from __future__ import annotations

import sys

from openpyxl import load_workbook

SECTIONS = ["القراءة", "الكتابة", "القواعد", "المفردات"]
HEADER = ["ID", "اسم الطالب", "#"] + list(range(1, 18)) + SECTIONS + ["المجموع"]
SHEETS = {"أول ثانوي 1": [101, 102, 103], "أول ثانوي 2": [201, 202]}
FORMULAS = {"U": "=SUM(N{r}:R{r})", "V": "=SUM(S{r}:T{r})", "W": "=SUM(I{r}:M{r})", "X": "=SUM(D{r}:H{r})", "Y": "=C{r}"}


def main() -> int:
    if len(sys.argv) < 2:
        print("usage: check_record_xlsx.py <record.xlsx>")
        return 2
    errs: list[str] = []
    wb = load_workbook(sys.argv[1])
    if wb.sheetnames != list(SHEETS):
        errs.append(f"sheet names {wb.sheetnames!r}")
    for name, ids in SHEETS.items():
        if name not in wb.sheetnames:
            continue
        ws = wb[name]
        if not ws.sheet_view.rightToLeft:
            errs.append(f"{name}: not right-to-left")
        head = [c.value for c in ws[1]]
        if head != HEADER:
            errs.append(f"{name}: header {head!r}")
        if not ws["A1"].font.b:
            errs.append(f"{name}: header not bold")
        for col, want in (("A", 10.63), ("B", 25.63), ("C", 5.63), ("D", 13.0), ("Y", 13.0)):
            got = ws.column_dimensions[col].width
            if got is None or abs(got - want) > 0.01:
                errs.append(f"{name}: width {col} {got} != {want}")
        for i, sid in enumerate(ids):
            r = i + 2
            if ws[f"A{r}"].value != sid:
                errs.append(f"{name}: A{r} {ws[f'A{r}'].value!r} != {sid}")
            if not isinstance(ws[f"B{r}"].value, str) or not ws[f"B{r}"].value:
                errs.append(f"{name}: B{r} has no name")
            for col, f in FORMULAS.items():
                want = f.format(r=r)
                if ws[f"{col}{r}"].value != want:
                    errs.append(f"{name}: {col}{r} {ws[f'{col}{r}'].value!r} != {want!r}")
        if ws.max_row != len(ids) + 1:
            errs.append(f"{name}: {ws.max_row} rows")
    first = wb["أول ثانوي 1"] if "أول ثانوي 1" in wb.sheetnames else None
    if first is not None:
        if first["C2"].value != 17 or first["D2"].value != 1:
            errs.append(f"C2/D2 {first['C2'].value!r} {first['D2'].value!r} (want numbers 17 and 1)")
        if first["C3"].value != 15 or first["D3"].value != 0:
            errs.append(f"C3/D3 {first['C3'].value!r} {first['D3'].value!r} (want 15 and 0)")
        if first["C4"].value is not None or first["D4"].value is not None:
            errs.append("row 4 (no scan) should be empty except ID, name and formulas")
    for e in errs:
        print(e)
    print("record sheet ok" if not errs else f"\n{len(errs)} problem(s)")
    return 1 if errs else 0


if __name__ == "__main__":
    sys.exit(main())
