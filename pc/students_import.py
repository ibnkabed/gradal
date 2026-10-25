"""Class lists from a file the teacher already has: Noor's Excel export of a class (xlsx, xls, or the HTML
table Noor sometimes saves as .xls), a CSV, or names pasted one per line.

    read_file(data, filename, class_name) -> {"classes": [{"name", "students": [{"name", "nationalId"}]}]}
    read_pasted(text, class_name)         -> the same shape
    then roster.merge(store, result)

The header row is found by its words, wherever it is: the name column («اسم الطالب», «الاسم», «Name»),
optionally the class («الفصل», «الشعبة») and the national ID («رقم الهوية», «السجل المدني»). Without a class
column every row goes to the class the teacher typed. Rows keep the file's order.
"""

from __future__ import annotations

import csv
import io
import re
from html.parser import HTMLParser

NAME_WORDS = ("اسم الطالب", "اسم الطالبة", "الاسم", "اسم", "name", "student")
CLASS_WORDS = ("الفصل", "الشعبة", "الصف", "class", "section")
ID_WORDS = ("رقم الهوية", "الهوية", "السجل المدني", "رقم الطالب", "national", "id")
DIGITS = str.maketrans("٠١٢٣٤٥٦٧٨٩", "0123456789")


class ImportError_(ValueError):
    """A message for the teacher, in Arabic."""


def _cell(v) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    return re.sub(r"\s+", " ", str(v)).strip()


class _Tables(HTMLParser):
    def __init__(self):
        super().__init__()
        self.rows, self.row, self.cell, self.inside = [], None, None, False

    def handle_starttag(self, tag, attrs):
        if tag == "tr":
            self.row = []
        elif tag in ("td", "th") and self.row is not None:
            self.cell = []

    def handle_endtag(self, tag):
        if tag in ("td", "th") and self.cell is not None and self.row is not None:
            self.row.append("".join(self.cell))
            self.cell = None
        elif tag == "tr" and self.row is not None:
            self.rows.append(self.row)
            self.row = None

    def handle_data(self, data):
        if self.cell is not None:
            self.cell.append(data)


def _rows(data: bytes, filename: str) -> list[list[str]]:
    name = (filename or "").lower()
    head = data[:512].lstrip().lower()
    if data[:2] == b"PK":
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        ws = wb.worksheets[0]
        return [[_cell(v) for v in row] for row in ws.iter_rows(values_only=True)]
    if data[:8] == b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1":
        import xlrd
        book = xlrd.open_workbook(file_contents=data)
        sh = book.sheet_by_index(0)
        return [[_cell(sh.cell_value(r, c)) for c in range(sh.ncols)] for r in range(sh.nrows)]
    text = None
    for enc in ("utf-8-sig", "cp1256", "utf-16"):
        try:
            text = data.decode(enc)
            break
        except UnicodeDecodeError:
            continue
    if text is None:
        raise ImportError_("تعذّرت قراءة الملف")
    if head.startswith((b"<", b"\xef\xbb\xbf<")) or "<table" in text[:5000].lower():
        p = _Tables()
        p.feed(text)
        return [[_cell(c) for c in r] for r in p.rows]
    if name.endswith((".csv", ".txt")) or "," in text or ";" in text or "\t" in text:
        dialect = csv.Sniffer().sniff(text[:4000], delimiters=",;\t") if text.strip() else csv.excel
        return [[_cell(c) for c in r] for r in csv.reader(io.StringIO(text), dialect)]
    return [[_cell(line)] for line in text.splitlines()]


def _find(header: list[str], words) -> int | None:
    low = [h.lower() for h in header]
    for w in words:
        for i, h in enumerate(low):
            if h == w or (len(w) > 3 and w in h):
                return i
    return None


def _is_name_header(cell: str) -> bool:
    """«اسم الطالب», «الاسم», «Student name»… but not a name that merely contains «اسم» (اسماعيل)."""
    t = re.sub(r"\s+", " ", cell.strip().lower())
    return t in NAME_WORDS or t.startswith(("اسم ال", "الاسم ", "student name", "name ")) or t in ("student name", "full name")


def _looks_like_name(s: str) -> bool:
    return bool(re.search(r"[A-Za-z؀-ۿ]", s)) and len(s.split()) >= 2 and not re.search(r"\d{3,}", s)


def from_rows(rows: list[list[str]], class_name: str) -> dict:
    rows = [r for r in rows if any(c for c in r)]
    header_at, name_col, class_col, id_col = None, None, None, None
    for i, r in enumerate(rows[:30]):
        c = next((k for k, cell in enumerate(r) if _is_name_header(cell)), None)
        if c is not None:
            header_at, name_col = i, c
            class_col, id_col = _find(r, CLASS_WORDS), _find(r, ID_WORDS)
            if id_col == name_col:
                id_col = None
            break
    if header_at is None:
        # no header: the column that holds the most names
        width = max((len(r) for r in rows), default=0)
        scores = [sum(1 for r in rows if c < len(r) and _looks_like_name(r[c])) for c in range(width)]
        if not scores or max(scores) == 0:
            raise ImportError_("لم أجد أسماء طلاب في الملف")
        name_col, header_at = scores.index(max(scores)), -1
    classes: dict[str, list] = {}
    for r in rows[header_at + 1:]:
        name = r[name_col] if name_col < len(r) else ""
        if not _looks_like_name(name):
            continue
        cname = (r[class_col] if class_col is not None and class_col < len(r) else "") or class_name
        cname = re.sub(r"\s+", " ", cname.translate(DIGITS)).strip()
        if not cname:
            raise ImportError_("اكتب اسم الفصل، فالملف لا يذكره")
        nid = (r[id_col] if id_col is not None and id_col < len(r) else "").translate(DIGITS)
        classes.setdefault(cname, []).append({"name": name, "nationalId": nid if re.fullmatch(r"\d{6,12}", nid) else ""})
    if not classes:
        raise ImportError_("لم أجد أسماء طلاب في الملف")
    return {"classes": [{"name": k, "students": v} for k, v in classes.items()]}


def read_file(data: bytes, filename: str, class_name: str = "") -> dict:
    try:
        rows = _rows(data, filename)
    except ImportError_:
        raise
    except Exception:  # noqa: BLE001
        raise ImportError_("تعذّرت قراءة الملف. احفظه من Excel بصيغة xlsx ثم أعد المحاولة.") from None
    return from_rows(rows, class_name)


def read_pasted(text: str, class_name: str) -> dict:
    names = [re.sub(r"\s+", " ", line).strip(" \t-•.0123456789٠١٢٣٤٥٦٧٨٩") for line in str(text or "").splitlines()]
    names = [n for n in names if n]
    if not names:
        raise ImportError_("الصق الأسماء، اسمًا في كل سطر")
    if not class_name.strip():
        raise ImportError_("اكتب اسم الفصل")
    return {"classes": [{"name": class_name.strip(), "students": [{"name": n, "nationalId": ""} for n in names]}]}
